"""Deterministic parsers for the V1 supported document formats."""

from __future__ import annotations

import io
import re
from abc import ABC, abstractmethod
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as DocxFile
from pypdf import PdfReader

from adaptive.ingestion.models import ParsedDocument, SourcePayload

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _decode_text(payload: bytes) -> str:
    return payload.decode("utf-8-sig", errors="replace").replace("\r\n", "\n").replace("\r", "\n")


class BaseParser(ABC):
    mime_types: frozenset[str] = frozenset()

    def supports(self, mime_type: str) -> bool:
        return mime_type.split(";", 1)[0].strip().lower() in self.mime_types

    @abstractmethod
    def parse(self, payload: bytes, mime_type: str | None = None) -> ParsedDocument:
        raise NotImplementedError


class PlainTextParser(BaseParser):
    mime_types = frozenset({"text/plain", "text/csv", "application/octet-stream"})

    def parse(self, payload: bytes, mime_type: str | None = None) -> ParsedDocument:
        return ParsedDocument(title=None, text=_decode_text(payload))


class MarkdownParser(BaseParser):
    mime_types = frozenset({"text/markdown", "text/x-markdown", "text/md"})

    def parse(self, payload: bytes, mime_type: str | None = None) -> ParsedDocument:
        text = _decode_text(payload)
        headings = [
            {"level": len(match.group(1)), "text": match.group(2).strip()}
            for match in re.finditer(r"(?m)^(#{1,6})[ \t]+(.+?)\s*$", text)
        ]
        title = headings[0]["text"] if headings and headings[0]["level"] == 1 else None
        return ParsedDocument(title=title, text=text, headings=headings)


class HtmlParser(BaseParser):
    mime_types = frozenset({"text/html", "application/xhtml+xml"})

    def parse(self, payload: bytes, mime_type: str | None = None) -> ParsedDocument:
        soup = BeautifulSoup(_decode_text(payload), "html.parser")
        for element in soup(["script", "style", "noscript", "template"]):
            element.decompose()
        title = soup.title.get_text(" ", strip=True) if soup.title else None
        headings = [
            {"level": int(element.name[1]), "text": element.get_text(" ", strip=True)}
            for element in soup.find_all(re.compile(r"^h[1-6]$"))
        ]
        tables = []
        for table_number, table in enumerate(soup.find_all("table")):
            rows = [
                [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                for row in table.find_all("tr")
            ]
            tables.append({"table_number": table_number, "rows": rows})
        return ParsedDocument(
            title=title,
            text=soup.get_text("\n", strip=True),
            headings=headings,
            tables=tables,
            metadata={
                "source_links": [anchor.get("href") for anchor in soup.find_all("a", href=True)]
            },
        )


class PdfParser(BaseParser):
    mime_types = frozenset({PDF_MIME})

    def parse(self, payload: bytes, mime_type: str | None = None) -> ParsedDocument:
        reader = PdfReader(io.BytesIO(payload))
        pages = []
        text_parts = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            pages.append({"page_number": page_number, "text": page_text})
            text_parts.append(page_text)
        metadata = {key.lstrip("/"): value for key, value in (reader.metadata or {}).items()}
        title = metadata.get("Title")
        return ParsedDocument(
            title=title, text="\n\n".join(text_parts), pages=pages, metadata=metadata
        )


class DocxParser(BaseParser):
    mime_types = frozenset({DOCX_MIME})

    def parse(self, payload: bytes, mime_type: str | None = None) -> ParsedDocument:
        document = DocxFile(io.BytesIO(payload))
        paragraphs: list[str] = []
        headings: list[dict[str, Any]] = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            paragraphs.append(value)
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                level_match = re.search(r"(\d+)$", paragraph.style.name)
                headings.append(
                    {"level": int(level_match.group(1)) if level_match else 1, "text": value}
                )
        tables = []
        for table_number, table in enumerate(document.tables):
            tables.append(
                {
                    "table_number": table_number,
                    "rows": [[cell.text.strip() for cell in row.cells] for row in table.rows],
                }
            )
        title = headings[0]["text"] if headings and headings[0]["level"] == 1 else None
        return ParsedDocument(
            title=title, text="\n\n".join(paragraphs), headings=headings, tables=tables
        )


PARSERS: tuple[BaseParser, ...] = (
    PdfParser(),
    DocxParser(),
    HtmlParser(),
    MarkdownParser(),
    PlainTextParser(),
)


def select_parser(mime_type: str) -> BaseParser:
    for parser in PARSERS:
        if parser.supports(mime_type):
            return parser
    raise ValueError(f"unsupported document MIME type: {mime_type}")  # noqa: TRY003


def parse_bytes(payload: bytes | SourcePayload, mime_type: str | None = None) -> ParsedDocument:
    source_metadata: dict[str, Any] = {}
    if isinstance(payload, SourcePayload):
        mime_type = payload.mime_type
        source_metadata = {
            **payload.metadata,
            "filename": payload.filename,
            **({"source_uri": payload.source_uri} if payload.source_uri else {}),
        }
        payload = payload.data
    if not mime_type:
        raise ValueError("mime_type is required")  # noqa: TRY003
    parsed = select_parser(mime_type).parse(payload, mime_type)
    if source_metadata:
        parsed = parsed.model_copy(update={"metadata": {**parsed.metadata, **source_metadata}})
    return parsed
